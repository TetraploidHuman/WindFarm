#!/usr/bin/env python3
"""Generate editable horizontal workflow PPT for three system layers."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "docs" / "figures" / "三层工作流图.pptx"

# Widescreen 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

C_TITLE = RGBColor(0x1A, 0x36, 0x5D)
C_SUB = RGBColor(0x4A, 0x55, 0x68)
C_LEGEND = RGBColor(0x71, 0x80, 0x96)
C_ARROW = RGBColor(0x4A, 0x55, 0x68)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_DARK = RGBColor(0x1A, 0x20, 0x2C)

# Layer palettes: (fill, border, accent for output)
PAL = {
    "phys": (RGBColor(0xE8, 0xF4, 0xFC), RGBColor(0x2B, 0x6C, 0xB0), RGBColor(0xE6, 0xF7, 0xED)),
    "cog": (RGBColor(0xE8, 0xF8, 0xF0), RGBColor(0x2F, 0x85, 0x5A), RGBColor(0xFE, 0xF3, 0xC7)),
    "plan": (RGBColor(0xFF, 0xF4, 0xE5), RGBColor(0xC2, 0x6E, 0x00), RGBColor(0xE8, 0xF0, 0xFE)),
}
C_INPUT = (RGBColor(0xF1, 0xF5, 0xF9), RGBColor(0x64, 0x74, 0x8B))


def _set_run(run, text: str, size: float, bold: bool = False, color: RGBColor = C_DARK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "微软雅黑"
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", "微软雅黑")


def _fill_text(shape, lines: list[tuple[str, float, bool]], align=PP_ALIGN.CENTER, color=C_DARK):
    tf = shape.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass
    for i, (text, size, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.text = text
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        for run in p.runs:
            _set_run(run, text, size, bold, color)


def rounded_box(slide, left, top, width, height, fill, border, lines, line_width=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(line_width)
    # softer corners
    try:
        adj = shape.adjustments
        adj[0] = 0.15
    except Exception:
        pass
    _fill_text(shape, lines, color=C_DARK)
    return shape


def arrow_right(slide, x1, y_mid, x2, color=C_ARROW):
    """Straight connector with arrow tip (line + triangle)."""
    # horizontal line
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y_mid, x2 - Inches(0.12), y_mid)
    conn.line.color.rgb = color
    conn.line.width = Pt(1.75)
    # arrow head
    head = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        x2 - Inches(0.18),
        y_mid - Inches(0.07),
        Inches(0.18),
        Inches(0.14),
    )
    head.fill.solid()
    head.fill.fore_color.rgb = color
    head.line.fill.background()
    return conn


def title_block(slide, title: str, subtitle: str):
    t = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.45))
    tf = t.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.clear()
    r = p.add_run()
    _set_run(r, title, 22, True, C_TITLE)

    s = slide.shapes.add_textbox(Inches(0.4), Inches(0.68), Inches(12.5), Inches(0.35))
    tf = s.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.clear()
    r = p.add_run()
    _set_run(r, subtitle, 12, False, C_SUB)


def legend(slide, text: str):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.clear()
    r = p.add_run()
    _set_run(r, text, 10, False, C_LEGEND)


def flow_boxes(slide, stages, y, h, gap, fill, border, out_fill=None):
    """Place equal-width stages horizontally and connect with arrows."""
    n = len(stages)
    margin_l = Inches(0.35)
    margin_r = Inches(0.35)
    usable = SLIDE_W - margin_l - margin_r
    # reserve gap space for arrows
    arrow_gap = gap
    box_w = (usable - arrow_gap * (n - 1)) / n
    shapes = []
    for i, stage in enumerate(stages):
        left = margin_l + i * (box_w + arrow_gap)
        f, b = fill, border
        if stage.get("kind") == "input":
            f, b = C_INPUT
        elif stage.get("kind") == "output" and out_fill is not None:
            f = out_fill
        lines = stage["lines"]
        sh = rounded_box(slide, left, y, box_w, h, f, b, lines)
        shapes.append((left, box_w, sh))
        if i > 0:
            prev_left, prev_w, _ = shapes[i - 1]
            x1 = prev_left + prev_w + Inches(0.02)
            x2 = left - Inches(0.02)
            arrow_right(slide, x1, y + h / 2, x2, border)
    return shapes


def slide_physics(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    fill, border, out = PAL["phys"]
    title_block(slide, "物理感知层工作流", "DEM物理下推 · 机载传感 · 残差特征")
    stages = [
        {
            "kind": "input",
            "lines": [
                ("输入数据", 11, True),
                ("SRTM1 DEM", 10, False),
                ("Open-Meteo 粗风", 10, False),
                ("10m / 100m", 9, False),
                ("空速/气压/GPS", 10, False),
            ],
        },
        {
            "lines": [
                ("地形特征提取", 11, True),
                ("坡度 · 坡向", 10, False),
                ("高差 Δh", 10, False),
                ("粗糙度 r", 10, False),
            ],
        },
        {
            "lines": [
                ("物理下推", 11, True),
                ("downscale_wind", 9, False),
                ("f_elev · f_slope", 10, False),
                ("对数剪切 · w", 10, False),
            ],
        },
        {
            "lines": [
                ("物理基线", 11, True),
                ("WindField", 9, False),
                ("u_phy, v_phy, w_phy", 9, False),
                ("形状 (Z,H,W)", 10, False),
            ],
        },
        {
            "lines": [
                ("特征构造", 11, True),
                ("62 维 FEATURE", 10, False),
                ("位置·地形·粗风", 9, False),
                ("物理基线·时序窗", 9, False),
            ],
        },
        {
            "kind": "output",
            "lines": [
                ("输出", 11, True),
                ("残差学习输入", 10, False),
                ("Observation", 10, False),
                ("局地风观测", 10, False),
            ],
        },
    ]
    flow_boxes(slide, stages, Inches(1.5), Inches(2.4), Inches(0.28), fill, border, out)
    # bottom note strip
    rounded_box(
        slide,
        Inches(0.5),
        Inches(4.4),
        Inches(12.3),
        Inches(2.2),
        RGBColor(0xF8, 0xFA, 0xFC),
        RGBColor(0xCB, 0xD5, 0xE1),
        [
            ("关键公式与产物（可编辑说明）", 12, True),
            ("f_elev = clip(1+0.0015·Δh, 0.7, 1.3)　　f_slope = clip(1+0.5·s·cosθ, 0.6, 1.4)", 11, False),
            ("垂直层：对数插值 10m↔100m 剪切；输出米级三维风场基线 + PhysicsDiagnostics", 11, False),
            ("下一步进入认知推理层：XGBoost 残差 Δu,Δv,Δw 与信念更新", 11, False),
        ],
        line_width=1.0,
    )
    legend(slide, "源码：physics.py / data_ingest.py / pipeline.py　　图3-1 分解 · 物理感知层")


def slide_cognition(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill, border, out = PAL["cog"]
    title_block(slide, "认知推理层工作流", "残差修正 · 三维信念地图 · 贝叶斯更新")
    stages = [
        {
            "kind": "input",
            "lines": [
                ("融合预报", 11, True),
                ("物理基线 u_phy", 10, False),
                ("XGBoost 残差", 10, False),
                ("Δu, Δv, Δw", 10, False),
                ("û = u_phy+Δu", 9, False),
            ],
        },
        {
            "lines": [
                ("先验推进", 11, True),
                ("apply_prediction", 9, False),
                ("过程噪声", 10, False),
                ("模式马尔可夫", 10, False),
                ("Kalman 融合", 10, False),
            ],
        },
        {
            "lines": [
                ("三维信念网格", 11, True),
                ("field_arrays", 9, False),
                ("wind / var / cov", 10, False),
                ("sink·neutral·uplift", 9, False),
            ],
        },
        {
            "lines": [
                ("观测校正", 11, True),
                ("update_with_observation", 8, False),
                ("空速+地速", 10, False),
                ("半径 R · 2×2 Kalman", 9, False),
                ("模式混合·平流", 10, False),
            ],
        },
        {
            "kind": "output",
            "lines": [
                ("规划用派生量", 11, True),
                ("uncertainty", 10, False),
                ("safety_penalty", 10, False),
                ("expected_energy_gain", 9, False),
            ],
        },
    ]
    flow_boxes(slide, stages, Inches(1.45), Inches(2.5), Inches(0.32), fill, border, out)

    # feedback dashed-like note
    fb = rounded_box(
        slide,
        Inches(3.5),
        Inches(4.2),
        Inches(6.3),
        Inches(0.55),
        RGBColor(0xEC, 0xFC, 0xF5),
        border,
        [("⟲ 边飞边认知：观测收缩不确定性 → 反馈下一时刻先验", 11, True)],
        line_width=1.25,
    )

    rounded_box(
        slide,
        Inches(0.5),
        Inches(5.0),
        Inches(12.3),
        Inches(1.7),
        RGBColor(0xF8, 0xFA, 0xFC),
        RGBColor(0xCB, 0xD5, 0xE1),
        [
            ("关键超参（config.json / BeliefConfig）", 12, True),
            ("process_noise=0.18　observation_noise=0.12　observation_radius=2　advection_gain=0.25　decay=0.04", 11, False),
            ("E = 0.55·max(w,0) − 0.85·max(−w,0) + 1.25·P_up − 0.95·P_sink − 0.28·S", 11, False),
        ],
        line_width=1.0,
    )
    legend(slide, "源码：belief.py / ml.py / pipeline.py　　图3-1 分解 · 认知推理层")


def slide_planning(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill, border, out = PAL["plan"]
    title_block(slide, "规划决策层工作流", "进度感知 · Anytime MPC · Dijkstra返航 · 滚动重规划")

    # Row 1: inputs -> parallel -> MPC
    y1 = Inches(1.35)
    h1 = Inches(1.55)
    # custom layout for parallel branch
    margin_l = Inches(0.35)
    w_in = Inches(2.1)
    w_par = Inches(2.35)
    w_mpc = Inches(2.6)
    w_pick = Inches(2.1)
    w_out = Inches(2.1)
    gap = Inches(0.22)

    x = margin_l
    rounded_box(
        slide,
        x,
        y1,
        w_in,
        Inches(2.35),
        C_INPUT[0],
        C_INPUT[1],
        [
            ("输入", 11, True),
            ("信念场 U/S/E", 10, False),
            ("DroneState", 10, False),
            ("Mission", 10, False),
            ("起终点·电池", 10, False),
        ],
    )
    x_prev = x + w_in
    x = x_prev + gap
    arrow_right(slide, x_prev + Inches(0.02), y1 + Inches(1.15), x - Inches(0.02), border)

    # parallel stack
    rounded_box(
        slide,
        x,
        y1,
        w_par,
        Inches(1.05),
        fill,
        border,
        [("Dijkstra 返航代价图", 10, True), ("compute_return_cost_map", 8, False)],
    )
    rounded_box(
        slide,
        x,
        y1 + Inches(1.2),
        w_par,
        Inches(1.05),
        fill,
        border,
        [("能量走廊引导线", 10, True), ("AGL / 侧向 corridor", 9, False)],
    )
    x_prev = x + w_par
    x = x_prev + gap
    arrow_right(slide, x_prev + Inches(0.02), y1 + Inches(1.15), x - Inches(0.02), border)

    rounded_box(
        slide,
        x,
        y1,
        w_mpc,
        Inches(2.35),
        fill,
        border,
        [
            ("Anytime 束搜索 MPC", 11, True),
            ("horizon=7  beam=28", 10, False),
            ("γ=0.93  rounds=4", 10, False),
            ("进度奖励 + 风险/安全", 9, False),
        ],
    )
    x_prev = x + w_mpc
    x = x_prev + gap
    arrow_right(slide, x_prev + Inches(0.02), y1 + Inches(1.15), x - Inches(0.02), border)

    rounded_box(
        slide,
        x,
        y1,
        w_pick,
        Inches(2.35),
        fill,
        border,
        [
            ("能量择优", 11, True),
            ("transition_energy_j", 9, False),
            ("MPC vs 直线AGL", 10, False),
            ("vs 走廊引导", 10, False),
        ],
    )
    x_prev = x + w_pick
    x = x_prev + gap
    arrow_right(slide, x_prev + Inches(0.02), y1 + Inches(1.15), x - Inches(0.02), border)

    rounded_box(
        slide,
        x,
        y1,
        w_out,
        Inches(2.35),
        out,
        border,
        [
            ("输出并执行", 11, True),
            ("航点/控制序列", 10, False),
            ("执行首步", 10, False),
            ("下发飞控", 10, False),
        ],
    )

    # replan loop bar
    rounded_box(
        slide,
        Inches(1.8),
        Inches(4.0),
        Inches(9.7),
        Inches(0.7),
        RGBColor(0xEF, 0xF6, 0xFF),
        RGBColor(0x3B, 0x82, 0xF6),
        [("⟲ 滚动重规划 NavigationEngine.step　→　预报 → 信念 → replan → 步进　（replan_interval_steps）", 11, True)],
        line_width=1.5,
    )

    rounded_box(
        slide,
        Inches(0.5),
        Inches(5.0),
        Inches(12.3),
        Inches(1.7),
        RGBColor(0xF8, 0xFA, 0xFC),
        RGBColor(0xCB, 0xD5, 0xE1),
        [
            ("关键代价与约束（可编辑）", 12, True),
            ("C = E_trans − R_progress + 22·λ_risk·U + 28·λ_safe·S + C_alt + C_vert", 11, False),
            ("risk_weight=1.2　safety_weight=1.5　reserve_energy_ratio=0.22　search_node_budget=9000", 11, False),
        ],
        line_width=1.0,
    )
    legend(slide, "源码：planner.py / controller.py / execution.py　　图3-1 分解 · 规划决策层")


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide_physics(prs)
    slide_cognition(prs)
    slide_planning(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
